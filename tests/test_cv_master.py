"""cv_master: Rendern mit Style-Pruefung wie v1, dazu Fassungen, Ergebniszeile und PDF."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import _common
import cv_master
import pytest
from docx import Document


def document_text(path: Path) -> str:
    """Absaetze und Tabellenzellen; der Name steht in der v1-Vorlage in einer Tabelle."""
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def test_load_data(dummy_data: Path):
    data = cv_master.load_data(dummy_data)
    assert data.name == "Anna Test"
    assert len(data.positions) == 2
    assert data.positions[0].bullets[0].startswith("Konzeption")
    assert data.education[0].detail == "Abschlussnote: 1,8"


def test_render_docx_fills_the_template(dummy_data: Path, cv_template: Path, tmp_path: Path):
    data = cv_master.load_data(dummy_data)
    out = tmp_path / "cv.docx"
    cv_master.render_docx(cv_template, data, out)
    text = document_text(out)
    assert "ANNA TEST" in text or "Anna Test" in text
    assert "Marketing Manager" in text
    assert "{{" not in text


def test_style_preservation_passes_for_a_clean_render(
    dummy_data: Path, cv_template: Path, tmp_path: Path
):
    data = cv_master.load_data(dummy_data)
    out = tmp_path / "cv.docx"
    cv_master.render_docx(cv_template, data, out)
    cv_master.validate_style_preservation(cv_template, out)


def _run(plugin_root: Path, *args: str, cwd: Path) -> tuple[int, dict, str]:
    proc = subprocess.run(
        [sys.executable, str(plugin_root / "scripts" / "cv_master.py"), *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        cwd=cwd,
    )
    lines = proc.stdout.rstrip("\n").splitlines()
    data = (
        json.loads(lines[-1][len(_common.RESULT_PREFIX) :])
        if lines and lines[-1].startswith(_common.RESULT_PREFIX)
        else {}
    )
    return proc.returncode, data, proc.stdout + proc.stderr


def test_cli_numbers_versions_and_reports_relative_paths(
    plugin_root: Path, dummy_data: Path, workdir: Path
):
    code, first, log = _run(
        plugin_root,
        "--data",
        str(dummy_data),
        "--name",
        "Anna Test",
        "--output-dir",
        "Bewerbungsmaterialien",
        cwd=workdir,
    )
    assert code == 0, log
    assert first["docx"] == "Bewerbungsmaterialien/Lebenslauf_Anna_Test_v1.docx"
    assert first["markdown"] == "Bewerbungsmaterialien/Lebenslauf_Anna_Test_v1.md"
    assert first["version"] == 1
    assert first["dateiname"] == "Lebenslauf_Anna_Test_v1.docx"
    assert first["pdf"] is None and first["style_drift"] is False
    assert (workdir / first["docx"]).is_file() and (workdir / first["markdown"]).is_file()

    code, second, log = _run(
        plugin_root,
        "--data",
        str(dummy_data),
        "--name",
        "Anna Test",
        "--output-dir",
        "Bewerbungsmaterialien",
        cwd=workdir,
    )
    assert code == 0, log
    assert second["version"] == 2
    assert (workdir / "Bewerbungsmaterialien" / "Lebenslauf_Anna_Test_v1.docx").is_file()

    code, third, log = _run(
        plugin_root,
        "--data",
        str(dummy_data),
        "--name",
        "Anna Test",
        "--output-dir",
        "Bewerbungsmaterialien",
        "--version",
        "7",
        cwd=workdir,
    )
    assert code == 0, log
    assert third["docx"].endswith("_v7.docx")


def test_cli_uses_the_plugin_template_by_default(
    plugin_root: Path, dummy_data: Path, workdir: Path
):
    code, data, log = _run(
        plugin_root,
        "--data",
        str(dummy_data),
        "--name",
        "Anna Test",
        "--output-dir",
        ".",
        cwd=workdir,
    )
    assert code == 0, log
    assert data["vorlage"].endswith("templates/Lebenslauf_template.docx")


def test_cli_missing_data_is_a_hard_error(plugin_root: Path, workdir: Path):
    code, data, log = _run(
        plugin_root, "--data", "fehlt.yml", "--name", "X", "--output-dir", ".", cwd=workdir
    )
    assert code == 1
    assert data == {}
    assert "FEHLER" in log


def test_cli_pdf_flag_with_libreoffice(
    plugin_root: Path, dummy_data: Path, workdir: Path, soffice: Path
):
    code, data, log = _run(
        plugin_root,
        "--data",
        str(dummy_data),
        "--name",
        "Anna Test",
        "--output-dir",
        "Bewerbungsmaterialien",
        "--pdf",
        cwd=workdir,
    )
    assert code == 0, log
    assert data["pdf"] == "Bewerbungsmaterialien/Lebenslauf_Anna_Test_v1.pdf"
    assert data["pdf_methode"] == "libreoffice"
    assert (workdir / data["pdf"]).is_file()


def test_cli_pdf_flag_without_converter_keeps_docx_and_exits_4(
    plugin_root: Path, dummy_data: Path, workdir: Path, monkeypatch: pytest.MonkeyPatch
):
    monkeypatch.setenv("JOBRADAR_SOFFICE", str(workdir / "fehlt.exe"))
    code, data, log = _run(
        plugin_root,
        "--data",
        str(dummy_data),
        "--name",
        "Anna Test",
        "--output-dir",
        ".",
        "--pdf",
        cwd=workdir,
    )
    if code == 0:
        pytest.skip("Word hat konvertiert")
    assert code == 4, log
    assert data["pdf"] is None
    assert data["hinweis"]
    assert "PDF" in log
    assert (workdir / data["docx"]).is_file()


def test_sanitize_filename_is_the_shared_one():
    assert cv_master.sanitize_filename is _common.sanitize_filename


def test_cli_top_level_list_yaml_is_a_hard_error(plugin_root: Path, workdir: Path):
    data_file = workdir / "liste.yml"
    data_file.write_text("- a\n- b\n", encoding="utf-8")
    code, data, log = _run(
        plugin_root, "--data", str(data_file), "--name", "A", "--output-dir", ".", cwd=workdir
    )
    assert code == 1, log
    assert data == {}
    assert "FEHLER" in log


def test_cli_output_dir_that_is_a_file_is_a_hard_error(
    plugin_root: Path, dummy_data: Path, workdir: Path
):
    (workdir / "Ordner").write_text("belegt", encoding="utf-8")
    code, data, log = _run(
        plugin_root,
        "--data",
        str(dummy_data),
        "--name",
        "A",
        "--output-dir",
        "Ordner",
        cwd=workdir,
    )
    assert code == 1, log
    assert data == {}
    assert "Zielordner nicht anlegbar" in log


def test_cli_template_that_is_no_docx_is_a_hard_error(
    plugin_root: Path, dummy_data: Path, workdir: Path
):
    code, data, log = _run(
        plugin_root,
        "--data",
        str(dummy_data),
        "--name",
        "A",
        "--output-dir",
        ".",
        "--template",
        str(dummy_data),
        cwd=workdir,
    )
    assert code == 1, log
    assert data == {}
    assert "Rendern fehlgeschlagen" in log


def test_cli_master_flag_switches_the_reported_art(
    plugin_root: Path, dummy_data: Path, workdir: Path
):
    """`/lebenslauf` rendert den Master mit `--master`; `/bewerbung` ohne die Fahne."""
    code, master, log = _run(
        plugin_root,
        "--data",
        str(dummy_data),
        "--name",
        "Anna Test",
        "--output-dir",
        "Bewerbungsmaterialien",
        "--master",
        cwd=workdir,
    )
    assert code == 0, log
    assert master["art"] == "master_lebenslauf"

    code, angepasst, log = _run(
        plugin_root,
        "--data",
        str(dummy_data),
        "--name",
        "Anna Test",
        "--output-dir",
        "Bewerbungen/Beispiel_GmbH",
        cwd=workdir,
    )
    assert code == 0, log
    assert angepasst["art"] == "lebenslauf"
