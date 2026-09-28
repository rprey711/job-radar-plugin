"""Rendering details of the CV: escaping, lists, empty sections, photo, metadata."""

from __future__ import annotations

import base64
import json
import re
import struct
import zipfile
import zlib
from datetime import datetime, timezone
from pathlib import Path

import _common
import cv_master
import pytest
import yaml
from docx import Document
from docx.shared import Mm

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "plugins" / "job-radar" / "templates" / "Lebenslauf_template.docx"
FIXTURES = Path(__file__).parent / "fixtures"


def _data(tmp_path: Path, **changes) -> Path:
    raw = yaml.safe_load((FIXTURES / "dummy_data.yml").read_text(encoding="utf-8"))
    raw.update(changes)
    path = tmp_path / "daten.yml"
    path.write_text(yaml.safe_dump(raw, allow_unicode=True), encoding="utf-8")
    return path


def _render(tmp_path: Path, **changes) -> Path:
    out = tmp_path / "cv.docx"
    cv_master.render_docx(TEMPLATE, cv_master.load_data(_data(tmp_path, **changes)), out)
    return out


def _xml(path: Path, part: str = "word/document.xml") -> str:
    with zipfile.ZipFile(path) as z:
        return z.read(part).decode("utf-8")


def test_ampersand_and_angle_brackets_survive(tmp_path: Path):
    out = _render(tmp_path, skills_section="Tools: A & B | Methoden: <C>")
    xml = _xml(out)
    assert "B.A. Marketing &amp; Kommunikation" in xml
    assert "TOOLS &amp; METHODEN" in xml
    assert "Tools: A &amp; B | Methoden: &lt;C&gt;" in xml


def _paragraphs(path: Path) -> list[str]:
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts += [p.text for p in cell.paragraphs]
    return texts


def test_empty_sections_disappear_with_their_heading(tmp_path: Path):
    texts = _paragraphs(_render(tmp_path, international=[], weiteres=[]))
    assert "INTERNATIONALER HINTERGRUND" not in texts
    assert "WEITERES" not in texts


def test_filled_sections_are_list_items_without_literal_bullets(tmp_path: Path):
    out = _render(
        tmp_path, international=["Auslandssemester Lyon"], weiteres=["Führerschein B", "Ehrenamt"]
    )
    doc = Document(str(out))
    items = [
        p
        for p in doc.paragraphs
        if p.text in ("Auslandssemester Lyon", "Führerschein B", "Ehrenamt")
    ]
    assert len(items) == 3
    for p in items:
        assert p._p.pPr is not None and p._p.pPr.numPr is not None
    assert not any(t.startswith("•") for t in _paragraphs(out))
    texts = _paragraphs(out)
    assert "INTERNATIONALER HINTERGRUND" in texts
    assert "WEITERES" in texts


def _list_items(path: Path) -> list[str]:
    return [
        p.text
        for p in Document(str(path)).paragraphs
        if p._p.pPr is not None and p._p.pPr.numPr is not None
    ]


def test_null_and_single_string_lists_render(tmp_path: Path):
    data = _data(tmp_path, weiteres=None, international="Auslandssemester Lyon")
    code = cv_master.main(
        ["--data", str(data), "--name", "Anna Test", "--output-dir", str(tmp_path / "out")]
    )
    assert code == 0
    out = tmp_path / "out" / "Lebenslauf_Anna_Test_v1.docx"
    items = _list_items(out)
    assert items.count("Auslandssemester Lyon") == 1
    assert "A" not in items
    texts = _paragraphs(out)
    assert "INTERNATIONALER HINTERGRUND" in texts
    assert "WEITERES" not in texts


def test_position_with_null_bullets_renders(tmp_path: Path):
    raw = yaml.safe_load((FIXTURES / "dummy_data.yml").read_text(encoding="utf-8"))
    positions = raw["positions"]
    positions[0]["bullets"] = None
    del positions[1]["bullets"]
    data = cv_master.load_data(_data(tmp_path, positions=positions))
    assert data.positions[0].bullets == [] and data.positions[1].bullets == []
    out = tmp_path / "cv.docx"
    cv_master.render_docx(TEMPLATE, data, out)
    texts = _paragraphs(out)
    assert any("Marketing Manager" in t for t in texts)
    assert not any(t.startswith("Konzeption") for t in texts)


def test_achievements_are_plain_list_items(tmp_path: Path):
    doc = Document(str(_render(tmp_path)))
    bullets = [p for p in doc.paragraphs if p.text.startswith("Aufbau einer Content-Strategie")]
    assert len(bullets) == 1
    p = bullets[0]
    assert p._p.pPr.numPr is not None
    for run in p.runs:
        assert not run.bold and not run.italic
        assert run.font.color is None or run.font.color.rgb is None


@pytest.mark.parametrize(
    "changes",
    [
        {},
        {"international": ["A1", "A2"], "weiteres": ["W1"], "geburtsdatum": "01.02.1990"},
    ],
    ids=["leer", "gefuellt"],
)
def test_no_empty_paragraphs_from_loop_tags(tmp_path: Path, changes: dict):
    out = _render(tmp_path, **changes)
    body = [p.text for p in Document(str(out)).paragraphs]
    # After the first heading only content lines, no blanks left by {% %} tags.
    start = body.index("BERUFSERFAHRUNG")
    assert "" not in body[start + 1 :]
    *text_cells, _photo_cell = _header_cells(out)
    for cell in text_cells:
        assert "" not in cell


def _header_cells(path: Path) -> list[list[str]]:
    row = Document(str(path)).tables[0].rows[0]
    return [[p.text for p in cell.paragraphs] for cell in row.cells]


def test_header_has_no_gap_without_birthdate(tmp_path: Path):
    *text_cells, _photo_cell = _header_cells(_render(tmp_path, geburtsdatum=""))
    for cell in text_cells:
        assert "" not in cell
    assert not any("Geboren" in t for cell in text_cells for t in cell)


def test_header_shows_the_birthdate_when_given(tmp_path: Path):
    cells = _header_cells(_render(tmp_path, geburtsdatum="01.02.1990"))
    assert "Geboren am 01.02.1990" in cells[0]
    for cell in cells[:-1]:
        assert "" not in cell


def test_blank_and_null_list_entries_are_dropped(tmp_path: Path):
    data = cv_master.load_data(
        _data(tmp_path, weiteres=[None, "Führerschein B", ""], international=["  ", None])
    )
    assert data.weiteres == ["Führerschein B"]
    assert data.international == []
    out = tmp_path / "cv.docx"
    cv_master.render_docx(TEMPLATE, data, out)
    items = _list_items(out)
    assert items.count("Führerschein B") == 1
    assert "" not in items
    texts = _paragraphs(out)
    assert "WEITERES" in texts
    assert "INTERNATIONALER HINTERGRUND" not in texts


def test_template_uses_calibri_and_has_no_picture_or_personal_metadata():
    with zipfile.ZipFile(TEMPLATE) as z:
        names = z.namelist()
        theme = z.read("word/theme/theme1.xml").decode("utf-8")
        props = z.read("docProps/core.xml").decode("utf-8") + z.read("docProps/app.xml").decode(
            "utf-8"
        )
        styles = z.read("word/styles.xml").decode("utf-8")
    assert not [n for n in names if n.startswith("word/media/")]
    assert 'typeface="Aptos' not in theme
    assert re.search(r'<a:latin typeface="Calibri"', theme)
    # Aptos's panose number next to the Calibri name would let Word substitute by panose.
    assert "02110004020202020204" not in theme
    assert re.search(r'<a:latin typeface="Calibri" panose="020F0502020204030204"', theme)
    defaults = re.search(r"<w:rPrDefault>.*?</w:rPrDefault>", styles, flags=re.S)
    assert defaults is not None
    assert 'w:ascii="Calibri"' in defaults.group(0)
    assert 'w:hAnsi="Calibri"' in defaults.group(0)
    assert "Raul" not in props and "Prey" not in props


def test_no_template_part_mentions_aptos():
    with zipfile.ZipFile(TEMPLATE) as z:
        parts = {n: z.read(n) for n in z.namelist() if n.endswith((".xml", ".rels"))}
    assert [n for n, data in parts.items() if b"Aptos" in data] == []
    fonts = parts["word/fontTable.xml"].decode("utf-8")
    calibri = re.search(r'<w:font w:name="Calibri">.*?</w:font>', fonts, flags=re.S)
    assert calibri is not None
    assert '<w:panose1 w:val="020F0502020204030204"/>' in calibri.group(0)


def test_template_default_language_is_german():
    with zipfile.ZipFile(TEMPLATE) as z:
        styles = z.read("word/styles.xml").decode("utf-8")
        settings = z.read("word/settings.xml").decode("utf-8")
    defaults = re.search(r"<w:rPrDefault>.*?</w:rPrDefault>", styles, flags=re.S)
    assert defaults is not None
    assert re.search(r'<w:lang w:val="de-DE"', defaults.group(0))
    assert '<w:themeFontLang w:val="de-DE"/>' in settings
    assert "en-DE" not in styles and "en-DE" not in settings


def test_template_metadata_carries_no_history():
    with zipfile.ZipFile(TEMPLATE) as z:
        core = z.read("docProps/core.xml").decode("utf-8")
        app = z.read("docProps/app.xml").decode("utf-8")
    assert "<cp:revision>1</cp:revision>" in core
    assert core.count("2026-01-01T00:00:00Z") == 2
    assert "2026-02-12" not in core and "2026-04-14" not in core
    assert "<TotalTime>0</TotalTime>" in app


def test_rendered_cv_carries_the_friends_name_as_author(tmp_path: Path):
    props = Document(str(_render(tmp_path))).core_properties
    assert props.author == "Anna Test"
    assert props.last_modified_by == "Anna Test"
    assert props.title == "Lebenslauf Anna Test"
    xml = _xml(tmp_path / "cv.docx", "docProps/core.xml")
    assert "Raul" not in xml


def test_rendered_cv_is_dated_at_render_time(tmp_path: Path):
    before = datetime.now(timezone.utc).replace(microsecond=0)
    out = _render(tmp_path)
    after = datetime.now(timezone.utc)
    props = Document(str(out)).core_properties
    assert before <= props.created <= after
    assert before <= props.modified <= after
    xml = _xml(out, "docProps/core.xml")
    assert "2013" not in xml and "2026-02-12" not in xml


PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _media(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        return [n for n in z.namelist() if n.startswith("word/media/")]


def test_photo_from_the_data_is_placed(tmp_path: Path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "Bewerbungsmaterialien").mkdir()
    (tmp_path / "Bewerbungsmaterialien" / "Bewerbungsfoto.png").write_bytes(PNG_1PX)
    out = tmp_path / "cv.docx"
    data = cv_master.load_data(_data(tmp_path, foto="Bewerbungsmaterialien/Bewerbungsfoto.png"))
    hinweis = cv_master.render_docx(TEMPLATE, data, out)
    assert hinweis is None
    assert len(_media(out)) == 1
    shapes = Document(str(out)).inline_shapes
    assert len(shapes) == 1
    assert shapes[0].width == Mm(32)


def test_no_photo_field_means_no_picture(tmp_path: Path):
    out = tmp_path / "cv.docx"
    hinweis = cv_master.render_docx(TEMPLATE, cv_master.load_data(_data(tmp_path)), out)
    assert hinweis is None
    assert _media(out) == []


def test_missing_photo_file_renders_without_and_says_so(tmp_path: Path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    out = tmp_path / "cv.docx"
    data = cv_master.load_data(_data(tmp_path, foto="Bewerbungsmaterialien/fehlt.jpg"))
    hinweis = cv_master.render_docx(TEMPLATE, data, out)
    assert _media(out) == []
    assert hinweis == "Foto nicht gefunden: Bewerbungsmaterialien/fehlt.jpg"


def test_cli_reports_a_missing_photo_in_the_result_line(tmp_path: Path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    data = _data(tmp_path, foto="Bewerbungsmaterialien/fehlt.jpg")
    code = cv_master.main(["--data", str(data), "--name", "Anna Test", "--output-dir", "out"])
    assert code == 0
    lines = capsys.readouterr().out.rstrip("\n").splitlines()
    assert "WARN Foto nicht gefunden: Bewerbungsmaterialien/fehlt.jpg" in lines
    result = json.loads(lines[-1][len(_common.RESULT_PREFIX) :])
    assert result["hinweis"] == "Foto nicht gefunden: Bewerbungsmaterialien/fehlt.jpg"
    assert result["docx"] == "out/Lebenslauf_Anna_Test_v1.docx"


HEIC_AS_JPG = b"\x00\x00\x00\x18ftypheic\x00\x00\x00\x00mif1heic" + b"\x00" * 200


def _png(width: int, height: int) -> bytes:
    """Graues RGB-PNG ohne Zusatzpaket."""
    raw = b"".join(b"\x00" + b"\x80\x80\x80" * width for _ in range(height))

    def chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


GREY_PNG = _png(60, 80)


@pytest.mark.parametrize(
    "name, content",
    [
        ("heic_als.jpg", HEIC_AS_JPG),
        ("abgeschnitten.png", GREY_PNG[:40]),
        ("abgeschnitten_mitte.png", GREY_PNG[: len(GREY_PNG) // 2]),
        ("leer.jpg", b""),
        ("abgeschnitten.gif", b"GIF89a"),
    ],
    ids=["heic", "png_kopf", "png_mitte", "leer", "gif_kopf"],
)
def test_unreadable_photo_renders_without_and_says_so(
    tmp_path: Path, monkeypatch, name: str, content: bytes
):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "Bewerbungsmaterialien").mkdir()
    (tmp_path / "Bewerbungsmaterialien" / name).write_bytes(content)
    out = tmp_path / "cv.docx"
    data = cv_master.load_data(_data(tmp_path, foto=f"Bewerbungsmaterialien/{name}"))
    hinweis = cv_master.render_docx(TEMPLATE, data, out)
    assert out.is_file()
    assert _media(out) == []
    assert hinweis is not None and hinweis.startswith("Foto nicht lesbar")
    assert f"Bewerbungsmaterialien/{name}" in hinweis


def test_render_failure_line_names_the_error(tmp_path: Path, monkeypatch, capsys):
    def kaputt(*_args, **_kwargs):
        raise ValueError()

    monkeypatch.setattr(cv_master, "render_docx", kaputt)
    data = _data(tmp_path)
    code = cv_master.main(
        ["--data", str(data), "--name", "Anna Test", "--output-dir", str(tmp_path / "out")]
    )
    assert code == 1
    line = next(x for x in capsys.readouterr().out.splitlines() if x.startswith("FEHLER: Rendern"))
    assert "ValueError" in line
    assert not line.rstrip().endswith(":")


def test_photo_and_pdf_hints_are_joined(tmp_path: Path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)

    def ohne_pdf(docx: Path, outdir: Path | None = None) -> dict:
        return {"docx": None, "pdf": None, "methode": None, "seiten": None, "hinweis": "kein PDF"}

    monkeypatch.setattr(cv_master.to_pdf, "to_pdf", ohne_pdf)
    data = _data(tmp_path, foto="Bewerbungsmaterialien/fehlt.jpg")
    code = cv_master.main(
        ["--data", str(data), "--name", "Anna Test", "--output-dir", "out", "--pdf"]
    )
    assert code == 4
    last = capsys.readouterr().out.rstrip("\n").splitlines()[-1]
    result = json.loads(last[len(_common.RESULT_PREFIX) :])
    assert result["hinweis"] == "Foto nicht gefunden: Bewerbungsmaterialien/fehlt.jpg; kein PDF"
