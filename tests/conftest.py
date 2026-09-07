"""Fixtures: Wegwerf-Ordner, Dummy-Daten, eine Marker-Vorlage fuer Anschreiben-Tests."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
PLUGIN = ROOT / "plugins" / "job-radar"
SCRIPTS = PLUGIN / "scripts"
TEMPLATES = PLUGIN / "templates"
FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture
def plugin_root() -> Path:
    return PLUGIN


@pytest.fixture
def cv_template() -> Path:
    return TEMPLATES / "Lebenslauf_template.docx"


@pytest.fixture
def dummy_data() -> Path:
    return FIXTURES / "dummy_data.yml"


@pytest.fixture
def cover_dummy_data() -> Path:
    return FIXTURES / "cover_dummy_data.yml"


def build_marker_cover_template(path: Path) -> None:
    """Minimale Anschreiben-Vorlage mit allen Markern; Anrede und Body mit expliziter Schrift,
    damit die Style-Pruefung etwas zu vergleichen hat."""
    doc = Document()
    doc.add_paragraph("{{ABSENDER}}")
    doc.add_paragraph("{{DATUM}}")
    doc.add_paragraph("{{EMPFAENGER}}")
    betreff = doc.add_paragraph()
    betreff.add_run("{{BETREFF}}").bold = True
    anrede = doc.add_paragraph("{{ANREDE}}")
    anrede.runs[0].font.name = "Calibri"
    anrede.runs[0].font.size = Pt(11)
    body = doc.add_paragraph("{{BODY}}")
    body.runs[0].font.name = "Calibri"
    body.runs[0].font.size = Pt(11)
    doc.add_paragraph("Mit freundlichen Grüßen")
    doc.add_paragraph("{{SIGNATUR}}")
    doc.save(str(path))


@pytest.fixture
def cover_template(tmp_path: Path) -> Path:
    path = tmp_path / "Anschreiben_template_test.docx"
    build_marker_cover_template(path)
    return path


@pytest.fixture
def workdir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """Ein leerer Arbeitsordner als cwd, wie der Job-Radar-Ordner beim Freund."""
    folder = tmp_path / "Job Radar"
    folder.mkdir()
    monkeypatch.chdir(folder)
    return folder


@pytest.fixture
def soffice() -> Path:
    """LibreOffice, sonst skip. Die CI installiert es, lokal ist es bei Raul da."""
    import _common

    found = _common.find_soffice()
    if found is None:
        pytest.skip("kein LibreOffice gefunden")
    return found


def copy_tree(src: Path, dst: Path) -> None:
    shutil.copytree(src, dst, dirs_exist_ok=True)
