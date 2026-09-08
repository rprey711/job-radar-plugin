"""read_docx: Absätze und Tabellen in Dokumentreihenfolge, Deckel, drei Fehlerwege."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import _common
import pytest
import read_docx
from docx import Document


def _docx(path: Path) -> Path:
    """Ein Dokument mit Absätzen, einem leeren Absatz, einer Tabelle und Umlauten."""
    doc = Document()
    doc.add_paragraph("Lebenslauf von Anna Müller")
    doc.add_paragraph("")
    doc.add_paragraph("Größte Stärke: Ausdauer")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Zeitraum"
    table.cell(0, 1).text = "Tätigkeit"
    table.cell(1, 0).text = "2020 bis 2024"
    table.cell(1, 1).text = "Projektleiterin"
    doc.add_paragraph("Schluss")
    doc.save(str(path))
    return path


def test_paragraphs_tables_and_umlauts_are_read(tmp_path: Path):
    text, ergebnis = read_docx.read_docx(_docx(tmp_path / "lebenslauf.docx"))
    assert "Lebenslauf von Anna Müller" in text
    assert "Größte Stärke: Ausdauer" in text
    assert "Zeitraum | Tätigkeit" in text
    assert "2020 bis 2024 | Projektleiterin" in text
    assert ergebnis["absaetze"] == 3
    assert ergebnis["tabellen"] == 1
    assert ergebnis["zeichen"] == len(text)
    assert ergebnis["gekuerzt"] is False
    assert ergebnis["datei"].endswith("lebenslauf.docx")


def test_document_order_is_kept(tmp_path: Path):
    text, _ = read_docx.read_docx(_docx(tmp_path / "reihenfolge.docx"))
    lines = text.splitlines()
    assert lines == [
        "Lebenslauf von Anna Müller",
        "Größte Stärke: Ausdauer",
        "Zeitraum | Tätigkeit",
        "2020 bis 2024 | Projektleiterin",
        "Schluss",
    ]


def test_max_zeichen_cuts_and_marks(tmp_path: Path):
    text, ergebnis = read_docx.read_docx(_docx(tmp_path / "kurz.docx"), max_zeichen=20)
    assert ergebnis["gekuerzt"] is True
    assert ergebnis["zeichen"] == 20
    assert text.startswith("Lebenslauf von Anna")
    assert text.endswith("[gekürzt]")


def test_missing_file_is_an_error(tmp_path: Path, capsys: pytest.CaptureFixture[str]):
    assert read_docx.main([str(tmp_path / "fehlt.docx")]) == 1
    assert "FEHLER" in capsys.readouterr().out


def test_a_file_that_is_no_docx_is_an_error(tmp_path: Path, capsys: pytest.CaptureFixture[str]):
    fake = tmp_path / "kein.docx"
    fake.write_text("Das ist nur Text, kein ZIP-Paket.", encoding="utf-8")
    assert read_docx.main([str(fake)]) == 1
    ausgabe = capsys.readouterr().out
    assert "FEHLER" in ausgabe and "DOCX" in ausgabe


def test_an_unreadable_file_is_an_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
):
    docx = _docx(tmp_path / "gesperrt.docx")

    def sperren(_path: str):
        raise OSError("Zugriff verweigert")

    monkeypatch.setattr(read_docx, "Document", sperren)
    assert read_docx.main([str(docx)]) == 1
    assert "FEHLER" in capsys.readouterr().out


def test_table_nested_in_a_cell_is_read(tmp_path: Path):
    """Eine Tabelle, die in einer Zelle einer anderen Tabelle steckt, muss mit auftauchen."""
    doc = Document()
    doc.add_paragraph("Start")
    aussen = doc.add_table(rows=1, cols=1)
    zelle = aussen.cell(0, 0)
    zelle.text = "Außenzelle"
    innen = zelle.add_table(rows=1, cols=2)
    innen.cell(0, 0).text = "Innen A"
    innen.cell(0, 1).text = "Innen B"
    doc.save(str(tmp_path / "verschachtelt.docx"))

    text, ergebnis = read_docx.read_docx(tmp_path / "verschachtelt.docx")
    assert "Innen A" in text
    assert "Innen B" in text
    assert ergebnis["tabellen"] == 2


def test_cli_prints_the_text_and_the_result_line(tmp_path: Path, plugin_root: Path):
    docx = _docx(tmp_path / "cli.docx")
    proc = subprocess.run(
        [sys.executable, str(plugin_root / "scripts" / "read_docx.py"), docx.name],
        capture_output=True,
        text=True,
        encoding="utf-8",
        cwd=tmp_path,
    )
    assert proc.returncode == 0, proc.stderr
    assert "Größte Stärke: Ausdauer" in proc.stdout
    line = proc.stdout.rstrip("\n").splitlines()[-1]
    data = json.loads(line[len(_common.RESULT_PREFIX) :])
    assert set(data) == {"datei", "zeichen", "absaetze", "tabellen", "gekuerzt"}
    assert data["datei"] == "cli.docx"
    assert data["absaetze"] == 3
    assert data["tabellen"] == 1
    assert data["gekuerzt"] is False
    assert data["zeichen"] > 0
