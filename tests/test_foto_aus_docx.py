"""foto_aus_docx: das größte Bild aus einem Word-Lebenslauf als Bewerbungsfoto ablegen."""

from __future__ import annotations

import base64
import json
import subprocess
import sys
from pathlib import Path

import _common
from docx import Document

PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _run(plugin_root: Path, *args: str, cwd: Path) -> tuple[int, dict]:
    proc = subprocess.run(
        [sys.executable, str(plugin_root / "scripts" / "foto_aus_docx.py"), *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        cwd=cwd,
    )
    last = proc.stdout.rstrip("\n").splitlines()[-1]
    return proc.returncode, json.loads(last[len(_common.RESULT_PREFIX) :])


def test_extracts_the_picture(plugin_root: Path, workdir: Path):
    png = workdir / "bild.png"
    png.write_bytes(PNG_1PX)
    doc = Document()
    doc.add_paragraph("Anna Test")
    doc.add_picture(str(png))
    (workdir / "Bewerbungsmaterialien").mkdir()
    doc.save(str(workdir / "Bewerbungsmaterialien" / "Lebenslauf_alt.docx"))
    code, result = _run(plugin_root, "Bewerbungsmaterialien/Lebenslauf_alt.docx", cwd=workdir)
    assert code == 0
    assert result["foto"] == "Bewerbungsmaterialien/Bewerbungsfoto.png"
    assert (workdir / "Bewerbungsmaterialien" / "Bewerbungsfoto.png").read_bytes() == PNG_1PX


def test_no_picture_is_a_clear_exit_1(plugin_root: Path, workdir: Path):
    doc = Document()
    doc.add_paragraph("ohne Bild")
    doc.save(str(workdir / "ohne.docx"))
    code, result = _run(plugin_root, "ohne.docx", cwd=workdir)
    assert code == 1
    assert result["foto"] is None
    assert result["hinweis"] == "Kein Bild im Dokument gefunden."


def test_not_a_docx_is_a_clear_exit_1(plugin_root: Path, workdir: Path):
    (workdir / "cv.pdf").write_bytes(b"%PDF-1.4")
    code, result = _run(plugin_root, "cv.pdf", cwd=workdir)
    assert code == 1
    assert result["foto"] is None
    assert "keine Word-Datei" in result["hinweis"]
