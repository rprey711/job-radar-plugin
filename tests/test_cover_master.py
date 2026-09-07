"""cover_master: Marker-Ersetzung wie v1, dazu Fassungen, Ergebniszeile, PDF und Seitenpruefung."""

from __future__ import annotations

import json
import subprocess
import sys
from datetime import date
from pathlib import Path

import _common
import cover_master
import pytest
from docx import Document


def test_format_date_de():
    assert cover_master.format_date_de(date(2026, 5, 21)) == "21. Mai 2026"
    assert cover_master.format_date_de(date(2026, 1, 3)) == "3. Januar 2026"


def test_render_replaces_all_markers(cover_template: Path, cover_dummy_data: Path, tmp_path: Path):
    data = cover_master.load_data(cover_dummy_data)
    out = tmp_path / "brief.docx"
    cover_master.render_docx(cover_template, data, out, name="Anna Test")
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "{{" not in text
    assert "Bewerbung als Marketing Manager" in text
    assert "Beispiel GmbH" in text
    assert text.count("Anna Test") == 2  # Absender und Signatur
    assert "Über ein persönliches Gespräch" in text


def test_expand_body_keeps_the_marker_style(cover_template: Path, tmp_path: Path):
    doc = Document(str(cover_template))
    cover_master.expand_body(doc, ["Erster Absatz.", "Zweiter Absatz."])
    paras = [p for p in doc.paragraphs if p.text.endswith("Absatz.")]
    assert len(paras) == 2
    assert all(p.runs[0].font.name == "Calibri" for p in paras)


def test_missing_body_marker_raises(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("{{ANREDE}}")
    path = tmp_path / "ohne_body.docx"
    doc.save(str(path))
    data = cover_master.CoverData(empfaenger="X", betreff="Y", anrede="Z", body=["a"])
    with pytest.raises(cover_master.MarkerNotFoundError):
        cover_master.render_docx(path, data, tmp_path / "out.docx")


def test_one_page_decision_prefers_the_pdf_page_count():
    short = cover_master.CoverData(empfaenger="X", betreff="Y", anrede="Z", body=["kurz"])
    assert cover_master.too_long(short, seiten=None, max_words=400) is False
    assert cover_master.too_long(short, seiten=2, max_words=400) is True
    long_body = cover_master.CoverData(
        empfaenger="X", betreff="Y", anrede="Z", body=["wort " * 500]
    )
    assert cover_master.too_long(long_body, seiten=None, max_words=400) is True
    assert cover_master.too_long(long_body, seiten=1, max_words=400) is False


def _run(plugin_root: Path, *args: str, cwd: Path) -> tuple[int, dict, str]:
    proc = subprocess.run(
        [sys.executable, str(plugin_root / "scripts" / "cover_master.py"), *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        cwd=cwd,
    )
    lines = proc.stdout.rstrip("\n").splitlines()
    data = (
        json.loads(lines[-1][len(_common.RESULT_PREFIX) :])
        if lines and lines[-1].startswith(_common.RESULT_PREFIX)
        else {}
    )
    return proc.returncode, data, proc.stdout + proc.stderr


def test_cli_names_versions_by_company(
    plugin_root: Path, cover_template: Path, cover_dummy_data: Path, workdir: Path
):
    args = (
        "--data",
        str(cover_dummy_data),
        "--name",
        "Anna Test",
        "--firma",
        "Beispiel GmbH",
        "--output-dir",
        "Bewerbungen/Beispiel_GmbH",
        "--template",
        str(cover_template),
    )
    code, first, log = _run(plugin_root, *args, cwd=workdir)
    assert code == 0, log
    assert first["docx"] == "Bewerbungen/Beispiel_GmbH/Anschreiben_Anna_Test_Beispiel_GmbH_v1.docx"
    assert first["art"] == "anschreiben" and first["version"] == 1 and first["zu_lang"] is False
    code, second, log = _run(plugin_root, *args, cwd=workdir)
    assert code == 0, log
    assert second["version"] == 2


def test_cli_default_template_is_the_plugin_one(
    plugin_root: Path, cover_dummy_data: Path, workdir: Path
):
    code, data, log = _run(
        plugin_root,
        "--data",
        str(cover_dummy_data),
        "--name",
        "Anna Test",
        "--firma",
        "Beispiel GmbH",
        "--output-dir",
        ".",
        cwd=workdir,
    )
    assert code == 0, log
    assert data["vorlage"].endswith("templates/Anschreiben_template.docx")


def test_cli_length_guard_exits_3_but_keeps_the_file(
    plugin_root: Path, cover_template: Path, workdir: Path
):
    data_file = workdir / "lang.yml"
    data_file.write_text(
        'empfaenger: |\n  Firma\n  Straße 1\nbetreff: B\nanrede: Hallo,\nbody:\n  - "'
        + "wort " * 450
        + '"\n',
        encoding="utf-8",
    )
    code, data, log = _run(
        plugin_root,
        "--data",
        str(data_file),
        "--name",
        "Anna Test",
        "--firma",
        "Firma",
        "--output-dir",
        ".",
        "--template",
        str(cover_template),
        cwd=workdir,
    )
    assert code == 3, log
    assert data["zu_lang"] is True
    assert (workdir / data["docx"]).is_file()


def test_cli_empty_body_is_a_hard_error(plugin_root: Path, cover_template: Path, workdir: Path):
    data_file = workdir / "leer.yml"
    data_file.write_text(
        "empfaenger: Firma\nbetreff: B\nanrede: Hallo,\nbody: []\n", encoding="utf-8"
    )
    code, data, log = _run(
        plugin_root,
        "--data",
        str(data_file),
        "--name",
        "A",
        "--firma",
        "F",
        "--output-dir",
        ".",
        "--template",
        str(cover_template),
        cwd=workdir,
    )
    assert code == 1 and data == {}
    assert "FEHLER" in log


def test_cli_pdf_page_count_decides(
    plugin_root: Path,
    cover_template: Path,
    cover_dummy_data: Path,
    workdir: Path,
    soffice: Path,
):
    code, data, log = _run(
        plugin_root,
        "--data",
        str(cover_dummy_data),
        "--name",
        "Anna Test",
        "--firma",
        "Beispiel GmbH",
        "--output-dir",
        ".",
        "--template",
        str(cover_template),
        "--pdf",
        cwd=workdir,
    )
    assert code == 0, log
    assert data["pdf"].endswith("_v1.pdf")
    assert data["seiten"] in (1, None)
