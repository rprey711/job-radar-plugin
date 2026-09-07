"""cover_master.py: Anschreiben-Filler des Job-Radar-Plugins.

Gießt den im Chat freigegebenen Brieftext in die Layout-Vorlage. Marker in der Vorlage, jeder
in einem eigenen Absatz: {{ABSENDER}} {{DATUM}} {{EMPFAENGER}} {{BETREFF}} {{ANREDE}} {{BODY}}
{{SIGNATUR}}. Der Body-Absatz wird pro Eintrag geklont, damit der Stil erhalten bleibt.

Aufruf:
    python cover_master.py --data <brief.yml> --name "<Name>" --firma "<Firma>"
                           --output-dir <ordner> [--template <vorlage.docx>] [--version N]
                           [--max-words 400] [--pdf]

Ausgabe: <ordner>/Anschreiben_<Name>_<Firma>_v<N>.docx (+ .pdf mit --pdf), letzte Zeile
JOBRADAR_RESULT {...}. Exit-Codes: 0 ok, 1 harter Fehler, 2 Style-Drift (Ausgabe bleibt),
3 länger als eine Seite (Ausgabe bleibt), 4 PDF nicht erzeugt (DOCX bleibt).
Exit 2 ohne Ergebniszeile ist ein Aufruffehler von argparse (falsche Argumente).

Einseitigkeit: mit PDF zählt die Seitenzahl des PDFs, ohne PDF die Wortzahl gegen --max-words.
"""

from __future__ import annotations

import argparse
import sys
from copy import deepcopy
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import _common
import to_pdf
import yaml
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


@dataclass
class CoverData:
    empfaenger: str
    betreff: str
    anrede: str
    body: list[str]
    datum: str = ""
    signatur: str = ""
    absender: str = ""


class StyleDriftError(Exception):
    """Output-Anschreiben weicht in Schriftart/Größe vom Template ab."""


class MarkerNotFoundError(Exception):
    """Ein Pflicht-Marker (z.B. {{BODY}}) fehlt im Template."""


def load_data(data_path: Path) -> CoverData:
    """Liest YAML, gibt CoverData zurück."""
    raw = yaml.safe_load(data_path.read_text(encoding="utf-8"))
    return CoverData(
        empfaenger=raw.get("empfaenger", ""),
        betreff=raw.get("betreff", ""),
        anrede=raw.get("anrede", ""),
        body=raw.get("body", []),
        datum=raw.get("datum", "") or "",
        signatur=raw.get("signatur", "") or "",
        absender=raw.get("absender", "") or "",
    )


_MONATE_DE = [
    "Januar",
    "Februar",
    "März",
    "April",
    "Mai",
    "Juni",
    "Juli",
    "August",
    "September",
    "Oktober",
    "November",
    "Dezember",
]


def format_date_de(d: date | None = None) -> str:
    """Datum im deutschen Format 'TAG. MONAT JAHR' ohne führende Null."""
    if d is None:
        d = date.today()
    return f"{d.day}. {_MONATE_DE[d.month - 1]} {d.year}"


def _set_single_run_text(paragraph, text: str) -> None:
    """Setze Text in den ersten Run, lösche restliche Runs (Style des ersten bleibt)."""
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    first = paragraph.runs[0]
    for r in list(paragraph.runs[1:]):
        r._element.getparent().remove(r._element)
    first.text = text


def replace_marker(doc, marker: str, value: str) -> bool:
    """Ersetze einen Einzelzeilen-Marker durch value. Style des Runs bleibt.

    Return True wenn der Marker gefunden + ersetzt wurde, sonst False (No-op).
    """
    for paragraph in doc.paragraphs:
        if marker in paragraph.text:
            for run in paragraph.runs:
                if marker in run.text:
                    run.text = run.text.replace(marker, value)
                    return True
            # Marker über mehrere Runs verteilt: ganzen Absatz neu setzen
            _set_single_run_text(paragraph, paragraph.text.replace(marker, value))
            return True
    return False


def expand_empfaenger(doc, lines: list[str]) -> bool:
    """Ersetze den {{EMPFAENGER}}-Marker durch mehrere Zeilen im selben Absatz.

    Zeilenumbrüche via w:br, damit der Empfängerblock eng bleibt (kein
    Absatz-Spacing zwischen den Zeilen). Style des Marker-Runs bleibt erhalten.
    Return True wenn Marker gefunden, sonst False (No-op).
    """
    lines = [ln for ln in lines if ln.strip()]
    if not lines:
        return False
    for paragraph in doc.paragraphs:
        if "{{EMPFAENGER}}" not in paragraph.text:
            continue
        # rPr des ersten Runs sichern (für Style-Kopie auf neue Runs)
        first = paragraph.runs[0] if paragraph.runs else paragraph.add_run("")
        rpr = first._element.find(qn("w:rPr"))
        # ersten Run mit erster Zeile belegen, restliche Runs entfernen
        for r in list(paragraph.runs[1:]):
            r._element.getparent().remove(r._element)
        first.text = lines[0]
        last = first
        for line in lines[1:]:
            last.add_break()  # Umbruch ans Ende des bisherigen Runs
            new_run = paragraph.add_run(line)
            if rpr is not None:
                new_run._element.insert(0, deepcopy(rpr))
            last = new_run
        return True
    return False


def expand_body(doc, paragraphs: list[str]) -> None:
    """Ersetze den {{BODY}}-Marker-Absatz durch N stilgleiche Absätze.

    Klont das <w:p>-Element des Markers pro Body-Eintrag (Style bleibt),
    fügt die Klone vor dem Marker ein, entfernt dann den Marker.
    Raised MarkerNotFoundError wenn kein {{BODY}}-Marker existiert.
    """
    marker_p = None
    for paragraph in doc.paragraphs:
        if "{{BODY}}" in paragraph.text:
            marker_p = paragraph
            break
    if marker_p is None:
        raise MarkerNotFoundError("{{BODY}} nicht im Template gefunden")

    for line in paragraphs:
        new_el = deepcopy(marker_p._element)
        new_para = Paragraph(new_el, marker_p._parent)
        _set_single_run_text(new_para, line)
        marker_p._element.addprevious(new_el)

    marker_p._element.getparent().remove(marker_p._element)


def render_docx(template_path: Path, data: CoverData, output_path: Path, name: str = "") -> None:
    """Rendert die Vorlage mit den Brief-Feldern, schreibt nach output_path."""
    doc = Document(str(template_path))

    datum = data.datum.strip() or format_date_de()
    signatur = data.signatur.strip() or name
    absender = data.absender.strip() or name

    replace_marker(doc, "{{BETREFF}}", data.betreff)
    replace_marker(doc, "{{ANREDE}}", data.anrede)
    replace_marker(doc, "{{DATUM}}", datum)
    # optionale Marker, No-op wenn nicht (mehr) vorhanden
    replace_marker(doc, "{{SIGNATUR}}", signatur)
    replace_marker(doc, "{{ABSENDER}}", absender)

    empf_lines = [ln for ln in data.empfaenger.splitlines() if ln.strip()]
    expand_empfaenger(doc, empf_lines)

    expand_body(doc, data.body)

    doc.save(str(output_path))


def _sample_run_style(doc, paragraph_index: int) -> dict:
    """Style des ersten Runs eines Absatzes (font_name/size/style_name)."""
    paragraphs = doc.paragraphs
    if paragraph_index < 0 or paragraph_index >= len(paragraphs):
        return {}
    para = paragraphs[paragraph_index]
    if not para.runs:
        return {
            "font_name": None,
            "font_size": None,
            "style_name": para.style.name if para.style else None,
        }
    run = para.runs[0]
    return {
        "font_name": run.font.name,
        "font_size": run.font.size,
        "style_name": para.style.name if para.style else None,
    }


def _find_marker_index(doc, marker: str) -> int | None:
    for i, p in enumerate(doc.paragraphs):
        if marker in p.text:
            return i
    return None


def validate_style_preservation(template_path: Path, output_path: Path) -> None:
    """Vergleicht Anrede- + Body-Absatz-Style zwischen Template und Output.

    Anrede + Body liegen vor den eingefügten Body-Absätzen bzw. an der Marker-
    Position, daher sind ihre Indizes zwischen Template und Output stabil.
    Bei Abweichung: StyleDriftError.
    """
    tdoc = Document(str(template_path))
    odoc = Document(str(output_path))

    drifts: list[str] = []
    for marker in ("{{ANREDE}}", "{{BODY}}"):
        idx = _find_marker_index(tdoc, marker)
        if idx is None:
            continue
        t_style = _sample_run_style(tdoc, idx)
        o_style = _sample_run_style(odoc, idx)
        for key in ("font_name", "font_size", "style_name"):
            if t_style.get(key) != o_style.get(key):
                drifts.append(
                    f"{marker}-{key}: template={t_style.get(key)} output={o_style.get(key)}"
                )

    if drifts:
        raise StyleDriftError("Style-Drift detected:\n  " + "\n  ".join(drifts))


def count_body_words(data: CoverData) -> int:
    """Anzahl Wörter im Body (alle Absätze zusammen)."""
    return sum(len(p.split()) for p in data.body)


def exceeds_one_page(data: CoverData, max_words: int = 400) -> bool:
    """Heuristik: überschreitet der Body das Ein-Seiten-Wort-Budget?"""
    return count_body_words(data) > max_words


def too_long(data: CoverData, seiten: int | None, max_words: int) -> bool:
    """Eine Seite ist die Grenze. Liegt ein PDF mit Seitenzahl vor, zählt sie; sonst die Wörter."""
    if seiten is not None:
        return seiten > 1
    return exceeds_one_page(data, max_words)


sanitize_filename = _common.sanitize_filename

DEFAULT_TEMPLATE = _common.PLUGIN_ROOT / "templates" / "Anschreiben_template.docx"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Anschreiben-Filler")
    parser.add_argument("--data", type=Path, required=True)
    parser.add_argument("--name", type=str, required=True)
    parser.add_argument("--firma", type=str, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument(
        "--version",
        type=int,
        default=None,
        help=(
            "Fassungsnummer, sonst die nächste freie, "
            "eine vorhandene Datei dieser Nummer wird ersetzt"
        ),
    )
    parser.add_argument("--max-words", type=int, default=400)
    parser.add_argument("--pdf", action="store_true")
    args = parser.parse_args(argv)

    if not args.template.is_file():
        print(f"FEHLER: Vorlage nicht gefunden: {args.template}")
        return 1
    if not args.data.is_file():
        print(f"FEHLER: Daten nicht gefunden: {args.data}")
        return 1
    if args.version is not None and args.version < 1:
        print("FEHLER: --version muss mindestens 1 sein")
        return 1
    try:
        data = load_data(args.data)
    except (AttributeError, TypeError, yaml.YAMLError) as exc:
        print(f"FEHLER: kein gültiges YAML: {exc}")
        return 1
    if not isinstance(data.body, list) or not all(isinstance(p, str) for p in data.body):
        print("FEHLER: body muss eine Liste von Absätzen sein, ein Text je Eintrag")
        return 1
    if not data.body:
        print("FEHLER: Kein Brieftext (body leer)")
        return 1
    if not isinstance(data.empfaenger, str):
        print("FEHLER: empfaenger muss Text sein, eine Zeile je Adresszeile")
        return 1
    if not data.empfaenger.strip():
        print("FEHLER: Kein Empfänger (empfaenger leer)")
        return 1

    out = args.output_dir
    try:
        out.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        print(f"FEHLER: Zielordner nicht anlegbar: {out} ({exc})")
        return 1
    stem = f"Anschreiben_{sanitize_filename(args.name)}_{sanitize_filename(args.firma)}"
    version = args.version or _common.next_version(out, stem)
    docx_path = out / f"{stem}_v{version}.docx"

    print(f"-> Rendere {docx_path.name} ...")
    try:
        render_docx(args.template, data, docx_path, name=args.name)
    except MarkerNotFoundError as exc:
        print(f"FEHLER: {exc}. Die Vorlage braucht einen {{{{BODY}}}}-Absatz.")
        return 1
    except Exception as exc:  # noqa: BLE001 - python-docx meldet Unlesbares auf viele Arten
        print(f"FEHLER: Vorlage nicht lesbar: {args.template} ({exc})")
        return 1

    drift = False
    try:
        validate_style_preservation(args.template, docx_path)
        print("OK Stil erhalten")
    except StyleDriftError as exc:
        print(f"WARN Style-Drift: {exc}")
        drift = True

    pdf_result = to_pdf.to_pdf(docx_path) if args.pdf else None
    if args.pdf:
        marke = "OK" if pdf_result["pdf"] else "WARN"
        print(f"{marke} PDF: {pdf_result['pdf'] or pdf_result['hinweis']}")
    seiten = pdf_result["seiten"] if pdf_result else None
    lang = too_long(data, seiten, args.max_words)
    if lang:
        grund = (
            f"{seiten} Seiten im PDF"
            if seiten
            else f"etwa {count_body_words(data)} Wörter, Budget {args.max_words}"
        )
        print(f"WARN Brief länger als eine Seite ({grund}), kürzen.")

    _common.print_result(
        {
            "art": "anschreiben",
            "docx": _common.relative_posix(docx_path),
            "pdf": pdf_result["pdf"] if pdf_result else None,
            "pdf_methode": pdf_result["methode"] if pdf_result else None,
            "seiten": seiten,
            "hinweis": pdf_result["hinweis"] if pdf_result else None,
            "woerter": count_body_words(data),
            "dateiname": docx_path.name,
            "version": version,
            "vorlage": args.template.resolve().as_posix(),
            "style_drift": drift,
            "zu_lang": lang,
        }
    )
    if drift:
        return 2
    if lang:
        return 3
    if args.pdf and not pdf_result["pdf"]:
        return 4
    return 0


if __name__ == "__main__":
    sys.exit(main())
