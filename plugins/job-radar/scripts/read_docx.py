"""Text eines DOCX ausgeben, damit die Abläufe vorhandene Unterlagen lesen können.

Aufruf: python read_docx.py <datei.docx> [--max-zeichen N]

Absätze und Tabellen kommen in Dokumentreihenfolge: das Skript läuft die Kinder des Body durch
(`w:p` und `w:tbl`), nicht erst `document.paragraphs` und dann `document.tables`. Leere Absätze
werden weggelassen, Tabellenzeilen als Zellen mit „ | “ dazwischen ausgegeben. `--max-zeichen`
schneidet den Text ab und hängt „[gekürzt]“ an; die Ergebniszeile sagt es mit `gekuerzt` auch.

Exit 0 mit dem Text, Exit 1 mit `FEHLER: …` bei fehlender, fremder oder unlesbarer Datei.
Letzte Zeile: JOBRADAR_RESULT mit `datei`, `zeichen` (Länge des ausgegebenen Textes ohne die
Kürzungsmarke), `absaetze` (die ausgegebenen, also nicht leeren), `tabellen`, `gekuerzt`.
"""

from __future__ import annotations

import argparse
import sys
from collections.abc import Iterator
from pathlib import Path

import _common
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ZELLTRENNER = " | "
KUERZUNGSMARKE = "[gekürzt]"


def _blocks(document) -> Iterator[Paragraph | Table]:
    """Absätze und Tabellen in der Reihenfolge, in der sie im Dokument stehen."""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def read_docx(datei: Path, max_zeichen: int | None = None) -> tuple[str, dict]:
    """Den Text lesen und beschreiben. Gibt den Text und die Ergebnisdaten zurück."""
    document = Document(str(datei))
    zeilen: list[str] = []
    absaetze = 0
    tabellen = 0
    for block in _blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                absaetze += 1
                zeilen.append(text)
        else:
            tabellen += 1
            for row in block.rows:
                zeilen.append(ZELLTRENNER.join(zelle.text.strip() for zelle in row.cells))
    text = "\n".join(zeilen)
    gekuerzt = max_zeichen is not None and len(text) > max_zeichen
    if gekuerzt:
        text = text[:max_zeichen]
    ergebnis = {
        "datei": _common.relative_posix(datei),
        "zeichen": len(text),
        "absaetze": absaetze,
        "tabellen": tabellen,
        "gekuerzt": gekuerzt,
    }
    if gekuerzt:
        text = f"{text}\n{KUERZUNGSMARKE}"
    return text, ergebnis


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Text eines DOCX ausgeben")
    parser.add_argument("datei", type=Path)
    parser.add_argument("--max-zeichen", type=int, default=None)
    args = parser.parse_args(argv)
    if not args.datei.is_file():
        print(f"FEHLER: Datei nicht gefunden: {args.datei}")
        return 1
    try:
        text, ergebnis = read_docx(args.datei, args.max_zeichen)
    except PackageNotFoundError:
        print(f"FEHLER: Keine lesbare DOCX-Datei: {args.datei}")
        return 1
    except OSError as fehler:
        print(f"FEHLER: Datei nicht lesbar: {args.datei} ({fehler})")
        return 1
    print(text)
    _common.print_result(ergebnis)
    return 0


if __name__ == "__main__":
    sys.exit(main())
